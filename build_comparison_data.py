"""Regenerate data/city_comparison.json from data/city_comparison.xlsx.

Usage (from the repository root):
    python build_comparison_data.py

The website reads the JSON at runtime; the Excel file ('city_comparison'
sheet) is the human-editable master (not published — see .gitignore). Edit
that file, run this script, then commit and push the JSON. Requires
openpyxl: pip install openpyxl

(This is the live pipeline. seed_comparison_data.py is the historical
one-time script that originally generated the xlsx — don't run it again,
it would overwrite your edits.)
"""
import json
import re
import sys

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("openpyxl is not installed. Run:  pip install openpyxl")

XLSX = "data/city_comparison.xlsx"
JSON_OUT = "data/city_comparison.json"
SHEET = "city_comparison"

COLS = ["city", "indicator_group", "indicator_name", "value", "display", "unit", "year",
        "source", "note", "severity", "display_order"]


def clean(value):
    """Normalise a cell value to a string, matching the old CSV output."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (int, float)):
        return str(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def main():
    wb = load_workbook(XLSX, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        sys.exit(f'Sheet "{SHEET}" not found in {XLSX}')
    ws = wb[SHEET]

    rows_out = []
    for row in ws.iter_rows(min_row=2, max_col=len(COLS), values_only=True):
        record = {col: clean(val) for col, val in zip(COLS, row)}
        if not record["city"] and not record["indicator_name"]:
            continue  # skip empty rows
        rows_out.append(record)

    with open(JSON_OUT, "w", encoding="utf-8") as f:
        json.dump(rows_out, f, ensure_ascii=False, indent=1)
        f.write("\n")

    print(f"Wrote {len(rows_out)} rows to {JSON_OUT}")


if __name__ == "__main__":
    main()
